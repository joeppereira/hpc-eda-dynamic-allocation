#!/usr/bin/env python3
"""
Generate editable Word document with data flow diagram
Creates a professional document with SmartArt-style flow diagram
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_border(cell, **kwargs):
    """Add borders to table cell"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    
    # Create border elements
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '12')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '0066CC')
        tcBorders.append(border)
    
    tcPr.append(tcBorders)

def create_dataflow_document():
    """Create Word document with editable data flow diagram"""
    
    doc = Document()
    
    # Set document margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    # Title
    title = doc.add_heading('ML-Driven Resource Optimization Flow', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Subtitle
    subtitle = doc.add_paragraph('Five-Stage Closed-Loop Process')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_format = subtitle.runs[0]
    subtitle_format.font.size = Pt(14)
    subtitle_format.font.color.rgb = RGBColor(102, 102, 102)
    
    doc.add_paragraph()  # Spacing
    
    # Create table for flow diagram (1 row, 5 columns)
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Define stages
    stages = [
        {
            'number': '1. SUBMIT',
            'component': 'User',
            'action': 'sbatch',
            'color': RGBColor(232, 244, 248)
        },
        {
            'number': '2. PREDICT',
            'component': 'ML API',
            'action': '(history)',
            'color': RGBColor(232, 244, 248)
        },
        {
            'number': '3. ALLOCATE',
            'component': 'SLURM',
            'action': '(schedule)',
            'color': RGBColor(232, 244, 248)
        },
        {
            'number': '4. EXECUTE',
            'component': 'Monitor',
            'action': '(actual)',
            'color': RGBColor(232, 244, 248)
        },
        {
            'number': '5. LEARN',
            'component': 'Update DB',
            'action': '(retrain)',
            'color': RGBColor(232, 244, 248)
        }
    ]
    
    # Fill table cells
    for idx, stage in enumerate(stages):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(1.8)
        
        # Add border
        add_border(cell)
        
        # Set background color
        shading_elm = OxmlElement('w:shd')
        shading_elm.set(qn('w:fill'), 'E8F4F8')
        cell._tc.get_or_add_tcPr().append(shading_elm)
        
        # Add content
        # Stage number
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run1 = p1.add_run(stage['number'])
        run1.font.bold = True
        run1.font.size = Pt(12)
        run1.font.color.rgb = RGBColor(0, 102, 204)
        
        # Component
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = p2.add_run(stage['component'])
        run2.font.size = Pt(11)
        run2.font.bold = True
        
        # Action
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run3 = p3.add_run(stage['action'])
        run3.font.size = Pt(10)
        run3.font.color.rgb = RGBColor(102, 102, 102)
    
    doc.add_paragraph()  # Spacing
    
    # Add arrow indicators
    arrow_para = doc.add_paragraph()
    arrow_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    arrow_run = arrow_para.add_run('→        →        →        →')
    arrow_run.font.size = Pt(24)
    arrow_run.font.color.rgb = RGBColor(0, 102, 204)
    arrow_run.font.bold = True
    
    doc.add_paragraph()  # Spacing
    
    # Feedback loop indicator
    feedback = doc.add_paragraph()
    feedback.alignment = WD_ALIGN_PARAGRAPH.CENTER
    feedback_run = feedback.add_run('↻ Continuous Improvement Loop ↻')
    feedback_run.font.size = Pt(12)
    feedback_run.font.color.rgb = RGBColor(0, 170, 0)
    feedback_run.font.bold = True
    
    doc.add_paragraph()  # Spacing
    doc.add_paragraph()  # Spacing
    
    # Detailed description
    doc.add_heading('Detailed Flow', level=2)
    
    descriptions = [
        ('Job Submission', 'User submits with design parameters → Job plugin intercepts'),
        ('Prediction', 'API queries database → ML model predicts CPU/memory/instance → Returns recommendations'),
        ('Resource Allocation', 'SLURM adjusts job resources → Selects optimal instance type → Dispatches to compute node'),
        ('Job Execution', 'Container runs workload → Monitors track CPU/memory/I/O → Job completes'),
        ('Learning Feedback', 'Collect actual metrics → Store in database → Retrain ML model → Improve next predictions')
    ]
    
    for idx, (title, desc) in enumerate(descriptions, 1):
        p = doc.add_paragraph()
        run_num = p.add_run(f'{idx}. ')
        run_num.font.bold = True
        run_num.font.color.rgb = RGBColor(0, 102, 204)
        
        run_title = p.add_run(f'{title}: ')
        run_title.font.bold = True
        
        run_desc = p.add_run(desc)
    
    doc.add_paragraph()  # Spacing
    
    # Key benefit
    benefit = doc.add_paragraph()
    benefit_run = benefit.add_run('Key Benefit: ')
    benefit_run.font.bold = True
    benefit_run.font.size = Pt(11)
    
    benefit_text = benefit.add_run(
        'Progressive elimination of resource waste and OOM failures through '
        'continuous learning from execution history.'
    )
    benefit_text.font.size = Pt(11)
    
    # Save document
    output_path = 'docs/diagrams/dataflow-diagram.docx'
    doc.save(output_path)
    
    print(f"✓ Word document created: {output_path}")
    print("✓ Fully editable - you can modify text, colors, and layout")
    print("✓ Professional format suitable for blog post or presentation")

if __name__ == "__main__":
    create_dataflow_document()
