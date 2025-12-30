#!/usr/bin/env python3
"""
Generate simple Word document matching the original ASCII diagram
Clean, minimal design
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_border(cell):
    """Add simple border to table cell"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '8')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tcBorders.append(border)
    
    tcPr.append(tcBorders)

def create_simple_dataflow():
    """Create simple Word document matching ASCII style"""
    
    doc = Document()
    
    # Title
    title = doc.add_heading('Data Flow Diagram', level=2)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    doc.add_paragraph()  # Spacing
    
    # Create simple table (2 rows, 5 columns)
    table = doc.add_table(rows=2, cols=5)
    table.style = 'Light Grid'
    
    # Row 1: Stage numbers and arrows
    stages = ['1. SUBMIT', '2. PREDICT', '3. ALLOCATE', '4. EXECUTE', '5. LEARN']
    
    for idx, stage in enumerate(stages):
        cell = table.rows[0].cells[idx]
        cell.width = Inches(1.5)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(stage)
        run.font.size = Pt(11)
        run.font.bold = True
        
        if idx < 4:
            # Add arrow
            arrow_p = cell.add_paragraph()
            arrow_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            arrow_run = arrow_p.add_run('──►')
            arrow_run.font.size = Pt(14)
    
    # Row 2: Components
    components = [
        ['User', 'sbatch'],
        ['ML API', '(history)'],
        ['SLURM', '(schedule)'],
        ['Monitor', '(actual)'],
        ['Update DB', '(retrain)']
    ]
    
    for idx, comp in enumerate(components):
        cell = table.rows[1].cells[idx]
        
        # First line
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run1 = p1.add_run(comp[0])
        run1.font.size = Pt(10)
        
        # Second line
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = p2.add_run(comp[1])
        run2.font.size = Pt(9)
        run2.font.color.rgb = RGBColor(102, 102, 102)
    
    doc.add_paragraph()  # Spacing
    
    # Feedback loop
    feedback = doc.add_paragraph()
    feedback.alignment = WD_ALIGN_PARAGRAPH.CENTER
    feedback_run = feedback.add_run('└──────────────┴───────────────┴───────────────┴───────────────┘')
    feedback_run.font.name = 'Courier New'
    feedback_run.font.size = Pt(10)
    
    loop_text = doc.add_paragraph()
    loop_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
    loop_run = loop_text.add_run('Continuous Improvement Loop')
    loop_run.font.size = Pt(10)
    loop_run.font.italic = True
    
    # Save document
    output_path = 'docs/diagrams/dataflow-simple.docx'
    doc.save(output_path)
    
    print(f"✓ Simple Word document created: {output_path}")
    print("✓ Matches original ASCII diagram style")
    print("✓ Fully editable in Microsoft Word")

if __name__ == "__main__":
    create_simple_dataflow()
