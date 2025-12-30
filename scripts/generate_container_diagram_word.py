#!/usr/bin/env python3
"""
Generate Word document for Singularity Container Execution diagram
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_shading(cell, color):
    """Add background color to cell"""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading_elm)

def create_container_diagram():
    """Create Word document for container execution flow"""
    
    doc = Document()
    
    # Title
    title = doc.add_heading('Singularity Container Execution Flow', level=2)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    doc.add_paragraph()
    
    # Main container box
    main_table = doc.add_table(rows=1, cols=1)
    main_table.style = 'Medium Grid 1 Accent 1'
    main_cell = main_table.rows[0].cells[0]
    
    # Title in main cell
    p_title = main_cell.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run('Compute Node: Singularity Container Execution')
    run_title.font.bold = True
    run_title.font.size = Pt(12)
    
    main_cell.add_paragraph()
    
    # Host layer
    p_host = main_cell.add_paragraph()
    p_host.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_host = p_host.add_run('Host: slurmd | SPANK Plugin | Fluent Bit')
    run_host.font.size = Pt(11)
    run_host.font.bold = True
    
    # Arrow
    p_arrow = main_cell.add_paragraph()
    p_arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_arrow = p_arrow.add_run('↓')
    run_arrow.font.size = Pt(16)
    
    # Singularity execution
    p_sing = main_cell.add_paragraph()
    p_sing.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_sing = p_sing.add_run('Singularity: singularity exec openroad.sif openroad ...')
    run_sing.font.size = Pt(10)
    run_sing.font.name = 'Courier New'
    
    main_cell.add_paragraph()
    
    # Container details - nested table
    container_table = main_cell.add_table(rows=1, cols=1)
    container_table.style = 'Light Grid Accent 1'
    container_cell = container_table.rows[0].cells[0]
    add_shading(container_cell, 'E8F4F8')
    
    # Container title
    p_cont_title = container_cell.paragraphs[0]
    p_cont_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_cont_title = p_cont_title.add_run('Container: OpenROAD EDA Tools')
    run_cont_title.font.bold = True
    run_cont_title.font.size = Pt(11)
    
    # Tools list
    p_tools = container_cell.add_paragraph()
    run_tools = p_tools.add_run('• Synthesis • Floorplan • Place • Route • Timing')
    run_tools.font.size = Pt(10)
    
    container_cell.add_paragraph()
    
    # Auto-mounted
    p_mount = container_cell.add_paragraph()
    run_mount = p_mount.add_run('Auto-mounted: /shared/designs, /shared/logs, $HOME')
    run_mount.font.size = Pt(9)
    run_mount.font.italic = True
    
    main_cell.add_paragraph()
    
    # Resources tracked
    p_resources = main_cell.add_paragraph()
    p_resources.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_resources = p_resources.add_run('Resources tracked: CPU, Memory, I/O (visible to host)')
    run_resources.font.size = Pt(10)
    run_resources.font.italic = True
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Why Singularity section
    doc.add_heading('Why Singularity?', level=3)
    
    benefits = [
        'No root required',
        'Auto-mounts filesystems',
        'Native SLURM integration',
        'Accurate monitoring',
        '<1% overhead'
    ]
    
    for benefit in benefits:
        p = doc.add_paragraph(benefit, style='List Bullet')
        p.runs[0].font.size = Pt(11)
    
    # Save
    output_path = 'docs/diagrams/container-execution.docx'
    doc.save(output_path)
    
    print(f"✓ Container execution diagram created: {output_path}")
    print("✓ Fully editable Word document")

if __name__ == "__main__":
    create_container_diagram()
